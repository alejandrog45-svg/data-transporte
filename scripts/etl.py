"""ETL: Excel (viajes 2025 + ventas/gastos 2026) -> data/data.json (publico, agregado).

Datos personales (nombre/telefono) NUNCA van al JSON publico. Los clientes
recurrentes se guardan en data/privado/ (ignorado por git) para uso interno.
Uso: E:\\python-portable\\python.exe scripts\\etl.py
"""
import json, re, hashlib, unicodedata, collections, csv, datetime as dt
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
VIAJES_XLSX = ROOT / "PLANILLA  MATRIZ  MARZO 2025.xlsx"
FIN_XLSX = ROOT / "AÑO 2026.xlsx"
OUT = ROOT / "data"
OUT.mkdir(exist_ok=True)
(OUT / "privado").mkdir(exist_ok=True)

ANIO_VIAJES = 2025
MESES = {"ENERO": 1, "FEB": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6, "JULIO": 7,
         "AGOSTO": 8, "SEPT": 9, "OCT": 10, "NOV": 11, "DIC": 12}
MES_NOMBRE = ["", "Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
DOW = ["Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
COMUNAS = ["Santiago Centro", "Las Condes", "Providencia", "Ñuñoa", "Vitacura", "Lo Barnechea", "La Reina",
           "Maipú", "Pudahuel", "Quilicura", "Colina", "Huechuraba", "Recoleta", "Independencia", "San Miguel",
           "La Florida", "Puente Alto", "San Bernardo", "Estación Central", "Peñalolén", "Macul", "Padre Hurtado",
           "Lo Prado", "La Pintana", "Chicureo", "Pirque", "Lampa", "La Granja", "Quinta Normal", "Cerrillos",
           "Renca", "Conchalí", "San Joaquín", "La Cisterna", "El Bosque", "Peñaflor", "Paine", "Valparaíso",
           "Viña del Mar", "Limache", "Casablanca"]


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


COM_N = [(norm(c), c) for c in sorted(COMUNAS, key=len, reverse=True)]


def sector(*textos):
    for t in textos:
        n = norm(t)
        if not n or n.startswith("AEROP"):
            continue
        for k, v in COM_N:
            if k in n:
                return v
        if re.search(r"\bSTGO\b|\bSANTIAGO\b", n):
            return "Santiago Centro"
        if "ESTACION" in n and "FERRO" in n:
            return "Estación Central"
    return "Sin identificar"


def medio_pago(v):
    n = norm(v)
    if not n:
        return "Sin dato"
    if "NO COBRAR" in n:
        return "No cobrado"
    if "EFECTIVO" in n:
        return "Efectivo"
    if "TARJETA" in n:
        return "Tarjeta"
    if "TRANSF" in n or "BANCO" in n or "DEPOSITO" in n:
        return "Transferencia"
    if "MERCADO" in n or n == "MP":
        return "Mercado Pago"
    return "Otro"


def conductor(v):
    n = norm(v)
    if not n or n in (".", "-"):
        return "Sin dato"
    if n.startswith("EXTERNO"):
        return "Externo"
    return n.title()


def fecha_hoja(titulo):
    n = norm(titulo)
    m = re.match(r"^(\d{1,2})\s*-?\s*([A-Z]+)", n)
    if not m:
        return None
    d, mes = int(m.group(1)), None
    for k, v in MESES.items():
        if m.group(2).startswith(k):
            mes = v
    if not mes:
        return None
    try:
        return dt.date(ANIO_VIAJES, mes, d)
    except ValueError:
        return None


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    return None


def viajes():
    wb = openpyxl.load_workbook(VIAJES_XLSX, read_only=True, data_only=True)
    calidad = collections.Counter()
    filas, vistos, dups = [], set(), []
    for ws in wb.worksheets:
        f = fecha_hoja(ws.title)
        if not f:
            calidad["hojas_ignoradas"] += 1
            continue
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not r or r[0] is None:
                continue
            r = tuple(r) + (None,) * 14
            clave = (f, tuple(str(x) for x in r[:13]))
            if clave in vistos:
                calidad["duplicados_eliminados"] += 1
                dups.append([f.isoformat(), ws.title] + [str(x) for x in r[:13]])
                continue
            vistos.add(clave)
            tarifa = num(r[8])
            if tarifa is None:
                calidad["sin_tarifa"] += 1
            hora = r[2].hour if hasattr(r[2], "hour") else None
            if hora is None:
                calidad["sin_hora"] += 1
            tel = re.sub(r"\D", "", str(r[1] or ""))[-9:]
            origen, destino = str(r[4] or ""), str(r[5] or "")
            sale_aero = norm(origen).startswith("AEROP")
            filas.append(dict(
                fecha=f, hora=hora, pax=num(r[3]) or 0, tarifa=tarifa or 0,
                conductor=conductor(r[9]), pago=medio_pago(r[10]),
                sentido="Desde aeropuerto" if sale_aero else ("Hacia aeropuerto" if norm(destino).startswith("AEROP") else "Otro"),
                sector=sector(destino if sale_aero else origen, origen, destino),
                tipo=str(r[7]).strip().title() if r[7] and str(r[7]).strip() not in (".", "") else "",
                cliente=hashlib.sha256(tel.encode()).hexdigest()[:8] if len(tel) >= 8 else None,
                nombre=str(r[0]).strip(), tel=tel))
    with open(OUT / "privado" / "duplicados_eliminados.csv", "w", newline="", encoding="utf-8-sig") as fh:
        csv.writer(fh).writerows([["fecha", "hoja", "nombre", "telefono", "hora", "pax", "origen", "destino", "vuelo", "tipo", "tarifa", "conductor", "pago", "equipaje", "arribo"]] + dups)
    return filas, calidad


def agg(filas, key):
    d = collections.defaultdict(lambda: [0, 0.0])
    for x in filas:
        d[key(x)][0] += 1
        d[key(x)][1] += x["tarifa"]
    return d


def tabla(d, n=None, orden="n"):
    rows = [dict(nombre=k, viajes=v[0], monto=round(v[1])) for k, v in d.items()]
    rows.sort(key=lambda r: -r["viajes" if orden == "n" else "monto"])
    return rows[:n] if n else rows


def finanzas():
    wb = openpyxl.load_workbook(FIN_XLSX, read_only=True, data_only=True)
    meses, dias, calidad = [], [], collections.Counter()
    for ws in wb.worksheets:
        n = norm(ws.title)
        mes = next((v for k, v in MESES.items() if k in n), None)
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        if not mes or len(rows) < 3:
            continue
        head = [norm(h) for h in rows[1]]
        cats = {i: re.sub(r"^PUBLICIDAD.*", "PUBLICIDAD", h) for i, h in enumerate(head) if i >= 2 and h}
        venta_m, gasto = 0.0, collections.Counter()
        for r in rows[2:]:
            d = num(r[0])
            if d is None or not 1 <= d <= 31:
                continue
            v = num(r[1]) if len(r) > 1 else None
            if len(r) > 1 and isinstance(r[1], str) and r[1].strip():
                calidad["texto_en_venta"] += 1
            g_dia = 0.0
            for i, c in cats.items():
                if i < len(r):
                    x = num(r[i])
                    if x is None and isinstance(r[i], str) and r[i].strip():
                        calidad["texto_en_gastos(xxx)"] += 1
                    if x:
                        gasto[c] += x
                        g_dia += x
            if v or g_dia:
                dias.append(dict(fecha=f"2026-{mes:02d}-{int(d):02d}", venta=round(v or 0), gasto=round(g_dia)))
            venta_m += v or 0
        total_g = sum(gasto.values())
        meses.append(dict(mes=mes, nombre=MES_NOMBRE[mes], venta=round(venta_m), gasto=round(total_g),
                          margen=round(venta_m - total_g),
                          gastos={k: round(v) for k, v in gasto.most_common()}))
    return meses, dias, calidad


def main():
    filas, cal = viajes()
    total = sum(x["tarifa"] for x in filas if x["pago"] != "No cobrado")
    cobrados = [x for x in filas if x["pago"] != "No cobrado" and x["tarifa"] > 0]
    por_mes = agg(filas, lambda x: x["fecha"].month)
    por_dia = agg(filas, lambda x: x["fecha"].isoformat())
    por_dow = agg(filas, lambda x: x["fecha"].weekday())
    heat = [[0] * 24 for _ in range(7)]
    for x in filas:
        if x["hora"] is not None:
            heat[x["fecha"].weekday()][x["hora"]] += 1
    cli = collections.defaultdict(lambda: dict(n=0, monto=0.0, nombre="", tel=""))
    for x in filas:
        if x["cliente"]:
            c = cli[x["cliente"]]
            c["n"] += 1; c["monto"] += x["tarifa"]; c["nombre"] = x["nombre"]; c["tel"] = x["tel"]
    recurrentes = sorted((c for c in cli.values() if c["n"] >= 2), key=lambda c: -c["n"])
    with open(OUT / "privado" / "clientes_recurrentes.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["nombre", "telefono", "viajes", "monto_clp"])
        for c in recurrentes:
            w.writerow([c["nombre"], c["tel"], c["n"], round(c["monto"])])
    meses_fin, dias_fin, cal_fin = finanzas()
    data = dict(
        meta=dict(generado=dt.datetime.now().isoformat(timespec="seconds"),
                  fuentes=["PLANILLA MATRIZ (viajes 2025)", "AÑO 2026 (ventas y gastos)"],
                  periodo_viajes=[min(por_dia), max(por_dia)],
                  calidad=dict(viajes=dict(cal), finanzas=dict(cal_fin),
                               sin_sector=sum(1 for x in filas if x["sector"] == "Sin identificar"),
                               pago_sin_dato=sum(1 for x in filas if x["pago"] == "Sin dato"))),
        viajes=dict(
            total=len(filas), ingresos=round(total), ticket_promedio=round(total / max(len(cobrados), 1)),
            no_cobrados=sum(1 for x in filas if x["pago"] == "No cobrado"),
            pasajeros=round(sum(x["pax"] for x in filas)),
            por_mes=[dict(mes=k, nombre=MES_NOMBRE[k], viajes=v[0], monto=round(v[1])) for k, v in sorted(por_mes.items())],
            por_dia=[dict(fecha=k, viajes=v[0], monto=round(v[1])) for k, v in sorted(por_dia.items())],
            por_dow=[dict(nombre=DOW[k], viajes=v[0], monto=round(v[1])) for k, v in sorted(por_dow.items())],
            heatmap=dict(dias=DOW, horas=list(range(24)), valores=heat),
            conductores=tabla(agg(filas, lambda x: x["conductor"]), 12, "monto"),
            pagos=tabla(agg(filas, lambda x: x["pago"])),
            sectores=tabla(agg(filas, lambda x: x["sector"]), 15),
            sentido=tabla(agg(filas, lambda x: x["sentido"])),
            clientes=dict(unicos=len(cli), recurrentes=len(recurrentes),
                          viajes_de_recurrentes=sum(c["n"] for c in recurrentes),
                          top_anonimo=[dict(id=f"Cliente {i+1:02d}", viajes=c["n"], monto=round(c["monto"]))
                                       for i, c in enumerate(recurrentes[:10])])),
        finanzas=dict(meses=meses_fin, dias=dias_fin))
    (OUT / "data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print("viajes:", len(filas), "| ingresos:", data["viajes"]["ingresos"], "| calidad:", data["meta"]["calidad"])
    print("meses fin:", [(m["nombre"], m["venta"], m["gasto"]) for m in meses_fin])


if __name__ == "__main__":
    main()
